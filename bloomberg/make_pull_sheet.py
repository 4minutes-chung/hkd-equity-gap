"""Build bloom_hkd_pull.xlsx from bloom_hkd_tickers.csv: one sheet per series with a BDH formula.

Open the xlsx on a Bloomberg PC with the Excel add-in; the formulas fill in. Save it as
data/raw/bloom_hkd_pull_filled.xlsx. Bloomberg data stays out of git (data/ is ignored).
"""
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).parent
tickers = pd.read_csv(HERE / "bloom_hkd_tickers.csv", dtype=str)

wb = openpyxl.Workbook()
index = wb.active
index.title = "index"
index.append(["sheet"] + list(tickers.columns))
for i, row in tickers.iterrows():
    sheet = f"{i + 1:02d}_{row['priority']}"
    index.append([sheet] + list(row))
    ws = wb.create_sheet(sheet)
    ws["A1"] = f"{row['series']} | {row['ticker']} | {row['field']}"
    if row["ticker"].startswith("unknown"):
        ws["A3"] = "Ticker unknown: search SECF, then put the BDH formula here"
    else:  # BDH fills dates in column A and values in column B from A3 down
        ws["A3"] = f'=BDH("{row["ticker"]}","{row["field"]}","{row["start"]}","{row["end"]}")'
wb.save(HERE / "bloom_hkd_pull.xlsx")
print(f"{len(tickers)} series -> {HERE / 'bloom_hkd_pull.xlsx'}")
